"""Anticipación de paradas: preview read-only del calendario del mes siguiente.

No escribe PMP, órdenes, programación ni backlog. La condición procede siempre
del maestro vigente; una OT ausente no significa que el plan sea incorrecto.
"""
from __future__ import annotations

from collections import Counter
import re
from datetime import date
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import text

from backend.database import get_engine
from backend.parsers.common import (
    header_mapping, normalize_text, workbook_from_bytes, is_operation_plan
)
from backend.services.v2_import_service import _parse_monthly


def _period(year: int, month: int) -> date:
    if not 2020 <= year <= 2100 or not 1 <= month <= 12:
        raise ValueError("Período inválido")
    return date(year, month, 1)


def _read_calendar(content: bytes) -> list[dict[str, Any]]:
    book = workbook_from_bytes(content)
    try:
        if not book.worksheets:
            raise ValueError("El archivo no contiene hojas de cálculo")
        mapping = header_mapping(book.worksheets[0], 1)
        if "ACTIVO" not in mapping or "PLANTRABAJO" not in mapping:
            raise ValueError(
                "La Lista de Calendario debe tener las columnas Activo y PlanTrabajo. "
                "Orden es opcional: se aceptan actividades sin OT."
            )
    finally:
        book.close()
    records = _parse_monthly(content)
    if not records:
        raise ValueError("El archivo no contiene actividades con Activo y PlanTrabajo")
    return records


def _plan_match_key(value: Any) -> str:
    """Ignore presentation spaces around dashes, but never erase words."""
    return re.sub(r"\\s*[-–—]\\s*", "-", normalize_text(value))


def classify_advance(rows: list[dict[str, Any]], plans: dict, assets: dict,
                     period: date) -> dict[str, Any]:
    result: list[dict[str, Any]] = []
    excluded = 0
    counts = Counter()
    no_ot = 0
    reasons = Counter()
    normalized_plans: dict[str, list[dict[str, Any]]] = {}
    for original_key, candidate in plans.items():
        normalized_plans.setdefault(_plan_match_key(original_key), []).append(candidate)
    for entry in rows:
        pkey = normalize_text(entry["plan_clave_software"])
        plan = plans.get(pkey)
        if plan is None:
            matches = normalized_plans.get(_plan_match_key(pkey), [])
            if len(matches) == 1:
                plan = matches[0]
        if (plan and plan["es_operacion"]) or (
            plan is None and is_operation_plan(pkey)
        ):
            excluded += 1
            continue
        asset = assets.get(normalize_text(entry["activo_codigo"]))
        stop = plan["tiempo_parada_efectivo_min"] if plan else None
        reason = ""
        if not plan:
            condition = "SIN DEFINIR"
            reason = "PLAN NO ENCONTRADO"
            observation = "El PlanTrabajo no coincide con un plan del maestro"
        elif stop is None:
            condition = "SIN DEFINIR"
            reason = "TIEMPO PARADA VACÍO"
            observation = "Plan encontrado, pero TiempoParada no fue definido en el software ni en la app"
        elif float(stop) > 0:
            condition = "EQUIPO DETENIDO"
            observation = ""
        elif float(stop) == 0:
            condition = "OPERANDO"
            observation = "TiempoParada = 0: no requiere parada"
        else:
            condition = "SIN DEFINIR"
            reason = "TIEMPO PARADA INVÁLIDO"
            observation = "TiempoParada negativo: corregir el maestro"
        if reason:
            reasons[reason] += 1
        if not asset:
            observation = (observation + "; " if observation else "") + "Equipo no encontrado en el maestro"
        ot = normalize_text(entry.get("numero_ot_raw"))
        ot = "" if ot in ("", "SIN ASIGNAR", "NONE", "NULL") else str(entry["numero_ot_raw"])
        if not ot:
            no_ot += 1
        people = plan["numero_personas_efectivo"] if plan else None
        time = entry.get("tiempo_planeado_min")
        if time is None and plan:
            time = plan["tiempo_ejecucion_min"]
        hh = (round(float(time) * float(people) / 60, 2)
              if time is not None and people is not None and float(people)>0 else None)
        record = {
            "fila_origen": entry["fila_origen"],
            "numero_ot": ot,
            "activo_codigo": entry["activo_codigo"],
            "activo_descripcion": (asset or {}).get("descripcion") or "",
            "area_codigo": (asset or {}).get("area_codigo") or "",
            "criticidad": (asset or {}).get("criticidad") or "",
            "plan_clave_software": entry["plan_clave_software"],
            "descripcion_plan": (plan or {}).get("descripcion_plan_trabajo") or "",
            "grupo": (plan or {}).get("grupo") or "",
            "especialidad": entry.get("especialidad") or (plan or {}).get("especialidad") or "SIN ESPECIALIDAD",
            "cronograma_planeacion": entry.get("cronograma_planeacion") or "",
            "tiempo_parada_min": float(stop) if stop is not None else None,
            "tiempo_ejecucion_min": float(time) if time is not None else None,
            "personas": float(people) if people is not None else None,
            "hh_estimadas": hh,
            "condicion": condition,
            "motivo_sin_definir": reason,
            "observacion": observation,
        }
        result.append(record)
        counts[condition] += 1
    result.sort(key=lambda r: (
        {"EQUIPO DETENIDO": 0, "OPERANDO": 1, "SIN DEFINIR": 2}[r["condicion"]],
        r["especialidad"], r["area_codigo"], r["activo_codigo"], r["plan_clave_software"],
        r["fila_origen"],
    ))
    return {
        "periodo": str(period),
        "total_archivo": len(rows),
        "excluidos_operacion": excluded,
        "total_mantenimiento": len(result),
        "equipo_detenido": counts["EQUIPO DETENIDO"],
        "operando": counts["OPERANDO"],
        "sin_definir": counts["SIN DEFINIR"],
        "sin_ot": no_ot,
        "sin_definir_por_plan": reasons["PLAN NO ENCONTRADO"],
        "sin_definir_por_tiempo": reasons["TIEMPO PARADA VACÍO"],
        "sin_definir_por_invalido": reasons["TIEMPO PARADA INVÁLIDO"],
        "filas": result,
        "criterios": {
            "parada": "TiempoParada efectivo > 0: EQUIPO DETENIDO; = 0: OPERANDO; ausente o plan sin maestro: SIN DEFINIR.",
            "ot": "Una fila sin número de OT también se analiza y se exporta.",
            "operacion": "Planes OPERACIÓN excluidos según el maestro vigente.",
            "periodo": "El periodo lo selecciona el usuario; la exportación provisional no cambia el PMP operativo.",
            "hh": "HH estimadas: tiempo planeado o de ejecución × número de personas / 60; no son horas reales.",
        },
    }


def preview_advance_stops(*, content: bytes, year: int, month: int) -> dict[str, Any]:
    period = _period(year, month)
    rows = _read_calendar(content)
    with get_engine().connect() as conn:
        plans = {
            normalize_text(f"{r['grupo']}-{r['plan_trabajo']}"): dict(r)
            for r in conn.execute(text("""
                SELECT grupo, plan_trabajo, es_operacion,
                       tiempo_parada_efectivo_min, numero_personas_efectivo,
                       tiempo_ejecucion_min, descripcion_plan_trabajo, especialidad
                FROM programacion.plan_trabajo
            """)).mappings()
        }
        assets = {
            normalize_text(r["codigo"]): dict(r)
            for r in conn.execute(text("""
                SELECT codigo, descripcion, area_codigo, criticidad
                FROM programacion.activo
            """)).mappings()
        }
    result = classify_advance(rows, plans, assets, period)
    # Sin alterar el maestro, mostrar cuándo el resultado OPERANDO = 0
    # obedece a que NO existe ningún tiempo de parada = 0 en la fuente.
    maintenance = [p for p in plans.values() if not p["es_operacion"]]
    result["maestro_planes"] = len(maintenance)
    result["maestro_planes_operando"] = sum(
        p["tiempo_parada_efectivo_min"] is not None
        and float(p["tiempo_parada_efectivo_min"]) == 0
        for p in maintenance
    )
    result["maestro_planes_sin_tiempo_parada"] = sum(
        p["tiempo_parada_efectivo_min"] is None for p in maintenance
    )
    return result


COLUMNS = [
    ("Condición", "condicion", 21),
    ("OT", "numero_ot", 20),
    ("Especialidad", "especialidad", 17),
    ("Área", "area_codigo", 16),
    ("Criticidad", "criticidad", 14),
    ("Equipo", "activo_codigo", 25),
    ("Descripción equipo", "activo_descripcion", 39),
    ("Plan de trabajo", "plan_clave_software", 54),
    ("Descripción plan", "descripcion_plan", 41),
    ("Grupo", "grupo", 14),
    ("Tiempo parada (min)", "tiempo_parada_min", 18),
    ("Tiempo ejecución (min)", "tiempo_ejecucion_min", 21),
    ("Personas", "personas", 13),
    ("HH estimadas", "hh_estimadas", 16),
    ("Cronograma planeación", "cronograma_planeacion", 30),
    ("Fila Excel original", "fila_origen", 20),
    ("Motivo SIN DEFINIR", "motivo_sin_definir", 30),
    ("Observación", "observacion", 54),
    ("Fecha propuesta parada", "fecha_propuesta", 23),
    ("Ventana o turno", "ventana_parada", 24),
    ("Responsable área", "responsable_area", 25),
    ("Estado coordinación", "estado_coordinacion", 28),
    ("Comentarios planeador", "comentarios_planeador", 44),
]


def export_advance_excel(result: dict[str, Any]) -> tuple[bytes, str]:
    book = Workbook()
    overview = book.active
    overview.title = "RESUMEN"
    overview.sheet_view.showGridLines = False
    navy, blue, pale, green, amber, grey = "17365D", "2F75B5", "E9F2FA", "0D7968", "CE8C22", "52667A"
    overview.merge_cells("A1:D2")
    overview["A1"] = "ANTICIPACIÓN DE PARADAS · TEAM FOODS – CEK"
    overview["A1"].font = Font(size=16,bold=True,color="FFFFFF")
    overview["A1"].fill = PatternFill("solid", fgColor=navy)
    overview["A1"].alignment = Alignment(vertical="center")
    overview["A4"],overview["B4"] = "PERÍODO PREPARADO", result["periodo"]
    for index,(label,key) in enumerate([
        ("ACTIVIDADES EN ARCHIVO","total_archivo"),
        ("MANTENIMIENTO","total_mantenimiento"),
        ("EQUIPO DETENIDO","equipo_detenido"),
        ("OPERANDO","operando"),
        ("SIN DEFINIR","sin_definir"),
        ("SIN OT ASIGNADA","sin_ot"),
        ("PLANES OPERACIÓN EXCLUIDOS","excluidos_operacion"),
    ],start=6):
        overview.cell(index,1,label)
        overview.cell(index,2,result[key])
        overview.cell(index,1).font=Font(bold=True,color=navy)
        overview.cell(index,2).font=Font(bold=True,color=green if key=="equipo_detenido" else navy,size=13)
        overview.cell(index,1).fill=PatternFill("solid",fgColor=pale)
    overview["A16"]="CRITERIO DE CONDICIÓN"
    overview["A16"].font=Font(bold=True,color=navy)
    overview["A17"]="TiempoParada > 0: EQUIPO DETENIDO | = 0: OPERANDO | sin dato: SIN DEFINIR."
    overview["A18"]="SIN DEFINIR no se asume OPERANDO ni DETENIDO: validar con el planeador."
    overview["A19"]="La ausencia de número de OT no impide la revisión anticipada."
    overview["A20"]="Este Excel NO altera PMP, programaciones, cierres ni backlog."
    overview["A21"]="HH son estimaciones del plan, no horas reales."
    overview["A22"]="Columnas amarillas: el planeador puede proponer fecha, ventana, área y estado de coordinación."
    overview["A24"]="DIAGNÓSTICO DE SIN DEFINIR"
    overview["A24"].font=Font(bold=True,color=navy,size=12)
    for n,label,key in [
        (25,"PLAN NO ENCONTRADO","sin_definir_por_plan"),
        (26,"TIEMPO PARADA VACÍO EN MAESTRO","sin_definir_por_tiempo"),
        (27,"TIEMPO PARADA INVÁLIDO","sin_definir_por_invalido"),
    ]:
        overview.cell(n,1,label)
        overview.cell(n,2,result.get(key,0))
        overview.cell(n,1).font=Font(bold=True,color=navy)
        overview.cell(n,1).fill=PatternFill("solid",fgColor=pale)
    overview.column_dimensions["A"].width=76
    overview.column_dimensions["B"].width=22
    overview.sheet_properties.pageSetUpPr.fitToPage=True

    tabs=[
        ("EQUIPO DETENIDO","EQUIPO DETENIDO","E5F2F0"),
        ("OPERANDO","OPERANDO","EDF4FA"),
        ("SIN DEFINIR","SIN DEFINIR","FFF3DE"),
    ]
    for name,condition,color in tabs:
        ws=book.create_sheet(name)
        ws.sheet_view.showGridLines=False
        ws.freeze_panes="F5"
        ws.sheet_properties.tabColor=(green if condition=="EQUIPO DETENIDO" else
                                     blue if condition=="OPERANDO" else amber)
        end_letter=get_column_letter(len(COLUMNS))
        ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=len(COLUMNS))
        title=ws.cell(1,1, f"{name} · ANTICIPACIÓN {result['periodo']}")
        title.font=Font(size=17,bold=True,color="FFFFFF")
        title.fill=PatternFill("solid",fgColor=navy)
        title.alignment=Alignment(vertical="center")
        ws.row_dimensions[1].height=32
        ws.merge_cells(start_row=2,start_column=1,end_row=2,end_column=len(COLUMNS))
        note=ws.cell(2,1,"Lista provisional para planeador: una fila por actividad del Excel, con o sin OT.")
        note.font=Font(size=10,italic=True,color=grey)
        ws.row_dimensions[2].height=24
        for i,(heading,key,width) in enumerate(COLUMNS,1):
            cell=ws.cell(4,i,heading)
            cell.font=Font(bold=True,color="FFFFFF",size=10)
            cell.fill=PatternFill("solid",fgColor=blue)
            cell.alignment=Alignment(wrap_text=True,vertical="center")
            ws.column_dimensions[get_column_letter(i)].width=width
        ws.row_dimensions[4].height=30
        subset=[r for r in result["filas"] if r["condicion"]==condition]
        for idx,entry in enumerate(subset,5):
            for col,(_,key,_) in enumerate(COLUMNS,1):
                value=entry.get(key)
                if key=="estado_coordinacion" and condition=="EQUIPO DETENIDO":
                    value="PENDIENTE DE COORDINAR"
                # Evitar que contenido importado del software ejecute fórmulas al abrir Excel.
                if isinstance(value,str) and value.lstrip().startswith(("=","+","-","@")):
                    value="'"+value
                cell=ws.cell(idx,col,value)
                cell.alignment=Alignment(vertical="top",wrap_text=key in (
                    "descripcion_plan","plan_clave_software","observacion"))
                if key in {"fecha_propuesta","ventana_parada","responsable_area",
                           "estado_coordinacion","comentarios_planeador"}:
                    cell.fill=PatternFill("solid",fgColor="FFF3D9")
                elif idx%2==0:
                    cell.fill=PatternFill("solid",fgColor=color)
                if key in ("hh_estimadas","personas","tiempo_parada_min","tiempo_ejecucion_min"):
                    cell.number_format="0.00"
            ws.row_dimensions[idx].height=23
        ws.auto_filter.ref=f"A4:{end_letter}{max(4,len(subset)+4)}"
        if condition=="EQUIPO DETENIDO" and subset:
            from openpyxl.worksheet.datavalidation import DataValidation
            coordination = DataValidation(
                type="list",
                formula1='"PENDIENTE DE COORDINAR,COORDINADA,REPROGRAMAR"',
                allow_blank=True,
            )
            ws.add_data_validation(coordination)
            status_column = next(i for i,(_,key,_) in enumerate(COLUMNS,1)
                                 if key=="estado_coordinacion")
            letter = get_column_letter(status_column)
            coordination.add(f"{letter}5:{letter}{len(subset)+4}")
        ws.page_setup.orientation="landscape"
        ws.sheet_properties.pageSetUpPr.fitToPage=True
    out=BytesIO()
    book.save(out)
    label=result["periodo"][:7].replace("-","")
    return out.getvalue(),f"anticipacion_paradas_{label}_team_foods_cek.xlsx"
