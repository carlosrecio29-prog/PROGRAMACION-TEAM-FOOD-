from __future__ import annotations

import base64
from collections import defaultdict
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from typing import Any

from sqlalchemy import text

from backend.database import get_engine
from backend.parsers.common import canonical_plan_name, is_operation_plan, normalize_text
from backend.services.v2_closure_service import (
    V2ClosureError, _index_calendar, _is_finalized, _parse_calendar, preview_week_closure,
)


class V2WeeklyTrackingError(ValueError):
    pass


def _pct(value: float, total: float) -> float:
    return round(value / total * 100.0, 1) if total else 0.0


def _week_number(week_start: date) -> int | None:
    first_day = week_start.replace(day=1)
    first_thursday_day = 1 + ((3 - first_day.weekday()) % 7)
    if week_start.day < first_thursday_day:
        return None
    return 1 + ((week_start.day - first_thursday_day) // 7)



def _ot_key(value: Any) -> str:
    return str(value or "").strip().upper()


def _finish_datetime(value: Any) -> datetime | None:
    raw = str(value or "").strip()
    if not raw:
        return None
    cleaned = raw.replace("Z", "+00:00")
    try:
        parsed = datetime.fromisoformat(cleaned)
        return parsed.replace(tzinfo=None)
    except ValueError:
        pass
    for fmt in ("%d/%m/%Y %H:%M:%S", "%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(raw, fmt)
        except ValueError:
            continue
    return None


def _match_backlog_calendar_item(
    item: dict[str, Any],
    by_ot: dict[str, list[dict[str, Any]]],
) -> tuple[dict[str, Any] | None, str]:
    """Busca una OT del backlog en todo el archivo; la fecha no participa."""
    ot = normalize_text(item.get("numero_ot"))
    if not ot or ot == "SIN ASIGNAR":
        return None, "SIN_NUMERO_OT"

    matches = by_ot.get(ot, [])
    if not matches:
        return None, "NO_ENCONTRADA"

    # La OT es la llave principal para backlog. Si solo aparece una vez,
    # su estado es suficiente, sin importar el rango de fechas descargado.
    if len(matches) == 1:
        return matches[0], "OT"

    # Si una OT agrupa varias actividades, usar equipo + plan para evitar
    # cerrar por error una actividad distinta.
    asset = normalize_text(item.get("activo_codigo"))
    plan = canonical_plan_name(item.get("plan_clave_software"))
    exact = [
        row for row in matches
        if normalize_text(row.get("activo")) == asset
        and canonical_plan_name(row.get("plan")) == plan
    ]
    if len(exact) == 1:
        return exact[0], "OT_EQUIPO_PLAN"

    # Cuando todas las filas de una OT múltiple están finalizadas, es seguro
    # considerar esa OT finalizada aunque el texto histórico del plan haya cambiado.
    if matches and all(_is_finalized(str(row.get("estado") or "")) for row in matches):
        return matches[0], "OT_MULTIPLE_TODAS_FINALIZADAS"

    return None, "OT_AMBIGUA"


def _reconcile_backlog_from_calendar(
    conn,
    *,
    specialty: str,
    calendar_rows: list[dict[str, Any]],
    recorded_by: str | None,
) -> dict[str, Any]:
    """Concilia backlog activo por OT/estado usando todo el Excel, sin filtro de fecha."""
    eligible_calendar_rows = []
    for row in calendar_rows:
        raw_plan = str(row.get("plan") or "")
        if is_operation_plan(raw_plan) or is_operation_plan(raw_plan.lstrip(" -–—:._/|")):
            continue
        eligible_calendar_rows.append(row)

    _, by_ot = _index_calendar(eligible_calendar_rows)
    specialty = _ot_key(specialty)

    backlog_rows = [dict(row) for row in conn.execute(text("""
        SELECT b.id AS backlog_id,b.orden_mantenimiento_id,b.especialidad,
               o.numero_ot,o.plan_clave_software,
               a.codigo AS activo_codigo,
               COALESCE(p.es_operacion,false) AS es_operacion
        FROM programacion.backlog_v2 b
        JOIN programacion.orden_mantenimiento o ON o.id=b.orden_mantenimiento_id
        JOIN programacion.activo a ON a.id=o.activo_id
        LEFT JOIN programacion.plan_trabajo p ON p.id=o.plan_trabajo_id
        WHERE b.estado_seguimiento<>'FINALIZADA'
          AND upper(COALESCE(b.especialidad,''))=:specialty
          AND COALESCE(p.es_operacion,false)=false
        ORDER BY b.id
    """), {"specialty": specialty}).mappings()]

    found = 0
    finalized_rows: list[tuple[dict[str, Any], dict[str, Any], str]] = []
    ambiguous = 0
    open_found = 0

    for item in backlog_rows:
        match, reason = _match_backlog_calendar_item(item, by_ot)
        if not match:
            if reason == "OT_AMBIGUA":
                ambiguous += 1
            continue
        found += 1
        if _is_finalized(str(match.get("estado") or "")):
            finalized_rows.append((item, match, reason))
        else:
            open_found += 1

    if finalized_rows:
        payload = [{
            "order_id": int(item["orden_mantenimiento_id"]),
            "state": normalize_text(match.get("estado")) or "FINALIZADO",
        } for item, match, _ in finalized_rows]

        conn.execute(text("""
            UPDATE programacion.orden_mantenimiento
            SET estado=:state,actualizado_en=now()
            WHERE id=:order_id
        """), payload)

        ids = [int(item["orden_mantenimiento_id"]) for item, _, _ in finalized_rows]
        conn.execute(text("""
            UPDATE programacion.backlog_v2
            SET estado_seguimiento='FINALIZADA',
                ultimo_resultado_cierre='FINALIZADA',
                finalizado_en=COALESCE(finalizado_en,now()),
                finalizado_por=COALESCE(:recorded_by,'SEGUIMIENTO LISTA DE CALENDARIO'),
                ultima_programacion_id=NULL,
                actualizado_en=now()
            WHERE orden_mantenimiento_id=ANY(CAST(:ids AS bigint[]))
              AND estado_seguimiento<>'FINALIZADA'
        """), {
            "ids": ids,
            "recorded_by": recorded_by,
        })

    return {
        "revisadas": len(backlog_rows),
        "encontradas": found,
        "finalizadas": len(finalized_rows),
        "abiertas": open_found,
        "ambiguas": ambiguous,
        "finalizadas_ots": [
            str(match.get("numero_ot") or item.get("numero_ot") or "").strip()
            for item, match, _ in finalized_rows
        ],
    }


def _detect_unplanned_rows(
    conn,
    *,
    programming: dict[str, Any],
    programmed_rows: list[dict[str, Any]],
    calendar_rows: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    """Detecta ejecución finalizada dentro del corte que no pertenece a la línea base."""
    programming_id = int(programming["id"])
    week_start = programming["semana_inicio"]
    week_end = programming["semana_fin"]
    specialty = _ot_key(programming.get("especialidad"))
    programmed_ots = {
        _ot_key(row.get("numero_ot"))
        for row in programmed_rows
        if _ot_key(row.get("numero_ot")) not in {"", "SIN ASIGNAR"}
    }

    candidates: dict[str, dict[str, Any]] = {}
    for row in calendar_rows:
        # La regla OPERACIÓN también aplica al seguimiento de ejecución real.
        # Se excluye por el texto del calendario antes de clasificar la OT.
        raw_plan = str(row.get("plan") or "")
        if is_operation_plan(raw_plan) or is_operation_plan(raw_plan.lstrip(" -–—:._/|")):
            continue
        ot = _ot_key(row.get("numero_ot"))
        if not ot or ot == "SIN ASIGNAR" or ot in programmed_ots:
            continue
        if not _is_finalized(str(row.get("estado") or "")):
            continue
        finished_at = _finish_datetime(row.get("fecha_fin_orden"))
        if not finished_at or not (week_start <= finished_at.date() <= week_end):
            continue
        row_specialty = _ot_key(row.get("especialidad"))
        if row_specialty and row_specialty != specialty:
            continue
        # Si la OT aparece repetida en el Excel, conservar la finalización más reciente.
        previous = candidates.get(ot)
        if previous is None or finished_at > previous["_finished_at"]:
            candidates[ot] = {**row, "_finished_at": finished_at}

    if not candidates:
        return []

    order_rows = [dict(row) for row in conn.execute(text("""
        SELECT o.id AS orden_mantenimiento_id,o.numero_ot,o.periodo,o.especialidad,
               o.tiempo_planeado_min,
               a.codigo AS activo_codigo,a.descripcion AS activo_descripcion,
               a.area_codigo,root.descripcion AS area_nombre,
               p.plan_trabajo,p.descripcion_grupo,p.numero_personas_efectivo,
               COALESCE(p.es_operacion,false) AS es_operacion,
               round(
                 COALESCE(o.tiempo_planeado_min,p.tiempo_ejecucion_min)/60.0
                 * COALESCE(p.numero_personas_efectivo,0),2
               ) AS hh_estimada
        FROM programacion.orden_mantenimiento o
        JOIN programacion.activo a ON a.id=o.activo_id
        LEFT JOIN programacion.activo root ON root.codigo='BA-'||a.area_codigo
        LEFT JOIN programacion.plan_trabajo p ON p.id=o.plan_trabajo_id
        WHERE o.numero_ot IS NOT NULL
          AND upper(btrim(o.numero_ot))=ANY(CAST(:ots AS text[]))
        ORDER BY o.periodo DESC,o.id DESC
    """), {"ots": list(candidates)}).mappings()]

    orders_by_ot: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for order in order_rows:
        orders_by_ot[_ot_key(order.get("numero_ot"))].append(order)

    resolved: list[tuple[str, dict[str, Any], dict[str, Any] | None]] = []
    order_ids: list[int] = []
    for ot, calendar in candidates.items():
        matches = orders_by_ot.get(ot, [])
        matching_specialty = [
            item for item in matches
            if not _ot_key(item.get("especialidad"))
            or _ot_key(item.get("especialidad")) == specialty
        ]
        order = (matching_specialty or matches or [None])[0]
        if order and _ot_key(order.get("especialidad")) not in {"", specialty}:
            continue
        # Si la OT histórica sí está enlazada al maestro, el maestro manda.
        if order and bool(order.get("es_operacion")):
            continue
        resolved.append((ot, calendar, order))
        if order:
            order_ids.append(int(order["orden_mantenimiento_id"]))

    other_programming: dict[int, dict[str, Any]] = {}
    backlog: set[int] = set()
    if order_ids:
        for row in conn.execute(text("""
            SELECT pi.orden_mantenimiento_id,ps.id AS programacion_origen_id,
                   ps.semana_inicio,ps.semana_fin,ps.estado
            FROM programacion.programacion_item_v2 pi
            JOIN programacion.programacion_semanal_v2 ps ON ps.id=pi.programacion_id
            WHERE pi.orden_mantenimiento_id=ANY(CAST(:ids AS bigint[]))
              AND ps.id<>:programming_id
            ORDER BY ps.semana_inicio DESC,ps.id DESC
        """), {"ids": order_ids, "programming_id": programming_id}).mappings():
            other_programming.setdefault(int(row["orden_mantenimiento_id"]), dict(row))

        backlog = {
            int(row[0])
            for row in conn.execute(text("""
                SELECT orden_mantenimiento_id
                FROM programacion.backlog_v2
                WHERE orden_mantenimiento_id=ANY(CAST(:ids AS bigint[]))
            """), {"ids": order_ids}).all()
        }

    result: list[dict[str, Any]] = []
    current_period = week_start.replace(day=1)
    for ot, calendar, order in resolved:
        order_id = int(order["orden_mantenimiento_id"]) if order else None
        previous_program = other_programming.get(order_id) if order_id else None
        if previous_program:
            origin = "PROGRAMADA_OTRA_SEMANA"
            if previous_program["semana_inicio"] > week_end:
                detail = (
                    "EJECUTADA ANTICIPADAMENTE · programada del "
                    f"{previous_program['semana_inicio']:%d/%m/%Y} al "
                    f"{previous_program['semana_fin']:%d/%m/%Y}"
                )
            elif previous_program["semana_fin"] < week_start:
                detail = (
                    "EJECUTADA FUERA DE SEMANA · programada del "
                    f"{previous_program['semana_inicio']:%d/%m/%Y} al "
                    f"{previous_program['semana_fin']:%d/%m/%Y}"
                )
            else:
                detail = "PROGRAMADA EN OTRA SEMANA"
        elif order_id and order_id in backlog:
            origin = "BACKLOG_NO_PROGRAMADO"
            detail = "BACKLOG EJECUTADO SIN ESTAR EN LA PROGRAMACIÓN DE ESTA SEMANA"
        elif order and order.get("periodo") == current_period:
            origin = "PMP_NO_PROGRAMADO"
            detail = "PMP DEL MES EJECUTADO SIN ESTAR EN LA PROGRAMACIÓN SEMANAL"
        else:
            origin = "EMERGENTE_NO_PLANIFICADA"
            detail = "OT FINALIZADA SIN PROGRAMACIÓN/PMP/BACKLOG IDENTIFICADO"

        raw_time = str(calendar.get("tiempo_planeado") or "").strip().replace(",", ".")
        try:
            raw_minutes = float(raw_time) if raw_time else None
        except ValueError:
            raw_minutes = None

        result.append({
            "programacion_id": programming_id,
            "orden_mantenimiento_id": order_id,
            "numero_ot": str(calendar.get("numero_ot") or ot).strip(),
            "activo_codigo": (order or {}).get("activo_codigo") or calendar.get("activo"),
            "activo_descripcion": (order or {}).get("activo_descripcion") or calendar.get("descripcion_activo"),
            "area_codigo": (order or {}).get("area_codigo"),
            "area_nombre": (order or {}).get("area_nombre"),
            "plan_trabajo": (order or {}).get("plan_trabajo") or calendar.get("plan"),
            "especialidad": (order or {}).get("especialidad") or calendar.get("especialidad") or specialty,
            "estado_calendario": calendar.get("estado"),
            "fecha_fin_orden": calendar["_finished_at"],
            "tiempo_planeado_min": (order or {}).get("tiempo_planeado_min") or raw_minutes,
            "hh_estimada": float((order or {}).get("hh_estimada") or 0),
            "origen": origin,
            "detalle_origen": detail,
            "programacion_origen_id": (
                int(previous_program["programacion_origen_id"]) if previous_program else None
            ),
            "semana_origen_inicio": previous_program["semana_inicio"] if previous_program else None,
            "semana_origen_fin": previous_program["semana_fin"] if previous_program else None,
        })
    return sorted(
        result,
        key=lambda row: (row["fecha_fin_orden"], _ot_key(row["numero_ot"])),
        reverse=True,
    )


def _area_summary(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    grouped: dict[str, dict[str, Any]] = defaultdict(lambda: {
        "area_codigo": "SIN ÁREA",
        "area_nombre": "",
        "programadas": 0,
        "finalizadas": 0,
        "pendientes": 0,
        "sin_coincidencia": 0,
        "hh_programadas": 0.0,
        "hh_finalizadas": 0.0,
        "hh_pendientes": 0.0,
    })
    for row in rows:
        key = str(row.get("area_codigo") or "SIN ÁREA")
        item = grouped[key]
        item["area_codigo"] = key
        item["area_nombre"] = row.get("area_nombre") or item["area_nombre"]
        hh = float(row.get("hh_programadas") or row.get("hh") or 0)
        item["programadas"] += 1
        item["hh_programadas"] += hh
        if row.get("finalizado") is True:
            item["finalizadas"] += 1
            item["hh_finalizadas"] += hh
        elif row.get("finalizado") is False:
            item["pendientes"] += 1
            item["hh_pendientes"] += hh
        else:
            item["sin_coincidencia"] += 1
    result = []
    for item in grouped.values():
        item["hh_programadas"] = round(item["hh_programadas"], 2)
        item["hh_finalizadas"] = round(item["hh_finalizadas"], 2)
        item["hh_pendientes"] = round(item["hh_pendientes"], 2)
        item["avance_ot_pct"] = _pct(item["finalizadas"], item["programadas"])
        item["avance_hh_pct"] = _pct(item["hh_finalizadas"], item["hh_programadas"])
        result.append(item)
    return sorted(result, key=lambda row: (-row["programadas"], row["area_codigo"]))


def _programmed_rows(conn, programming_id: int) -> list[dict[str, Any]]:
    return [dict(row) for row in conn.execute(text("""
        SELECT pi.id AS programacion_item_id,pi.orden_mantenimiento_id,o.numero_ot,
               a.area_codigo,root.descripcion AS area_nombre,pi.hh_programadas,
               NULL::text AS estado_calendario,NULL::boolean AS finalizado,
               'SIN SEGUIMIENTO'::text AS coincidencia,
               a.codigo AS activo_codigo,a.descripcion AS activo_descripcion,
               p.plan_trabajo,p.descripcion_grupo,
               pi.requiere_parada,pi.origen,pi.origen_backlog,
               COALESCE(o.tiempo_planeado_min,p.tiempo_ejecucion_min) AS tiempo_min,
               p.numero_personas_efectivo AS personas
        FROM programacion.programacion_item_v2 pi
        JOIN programacion.orden_mantenimiento o ON o.id=pi.orden_mantenimiento_id
        JOIN programacion.activo a ON a.id=o.activo_id
        LEFT JOIN programacion.activo root ON root.codigo='BA-'||a.area_codigo
        LEFT JOIN programacion.plan_trabajo p ON p.id=o.plan_trabajo_id
        WHERE pi.programacion_id=:id
        ORDER BY a.area_codigo,o.numero_ot NULLS LAST,a.codigo
    """), {"id": programming_id}).mappings()]


def get_week_tracking(programming_id: int) -> dict[str, Any]:
    with get_engine().connect() as conn:
        programming = conn.execute(text("""
            SELECT id,semana_inicio,semana_fin,especialidad,estado
            FROM programacion.programacion_semanal_v2
            WHERE id=:id
        """), {"id": programming_id}).mappings().first()
        if not programming:
            raise V2WeeklyTrackingError("Programación semanal no encontrada")

        history = [dict(row) for row in conn.execute(text("""
            SELECT id,archivo_nombre,registrado_por,registrado_en,total_items,
                   finalizados,pendientes,sin_coincidencia,
                   hh_programadas,hh_finalizadas,hh_pendientes,
                   avance_ot_pct,avance_hh_pct
            FROM programacion.seguimiento_semanal_v2
            WHERE programacion_id=:id
            ORDER BY registrado_en DESC,id DESC
            LIMIT 30
        """), {"id": programming_id}).mappings()]

        latest_id = int(history[0]["id"]) if history else None
        if latest_id:
            latest_rows = [dict(row) for row in conn.execute(text("""
                SELECT s.programacion_item_id,s.orden_mantenimiento_id,s.numero_ot,
                       s.area_codigo,s.area_nombre,s.hh_programadas,
                       s.estado_calendario,s.finalizado,s.coincidencia,
                       a.codigo AS activo_codigo,a.descripcion AS activo_descripcion,
                       p.plan_trabajo,p.descripcion_grupo,
                       pi.requiere_parada,pi.origen,pi.origen_backlog,
                       COALESCE(o.tiempo_planeado_min,p.tiempo_ejecucion_min) AS tiempo_min,
                       p.numero_personas_efectivo AS personas
                FROM programacion.seguimiento_semanal_item_v2 s
                JOIN programacion.programacion_item_v2 pi ON pi.id=s.programacion_item_id
                JOIN programacion.orden_mantenimiento o ON o.id=s.orden_mantenimiento_id
                JOIN programacion.activo a ON a.id=o.activo_id
                LEFT JOIN programacion.plan_trabajo p ON p.id=o.plan_trabajo_id
                WHERE s.seguimiento_id=:tracking_id
                ORDER BY s.area_codigo,s.finalizado DESC NULLS LAST,s.numero_ot NULLS LAST,a.codigo
            """), {"tracking_id": latest_id}).mappings()]
        else:
            latest_rows = _programmed_rows(conn, programming_id)

        unplanned_rows = [dict(row) for row in conn.execute(text("""
            SELECT n.id,n.orden_mantenimiento_id,n.numero_ot,n.activo_codigo,
                   n.activo_descripcion,n.area_codigo,n.area_nombre,n.plan_trabajo,
                   n.especialidad,n.estado_calendario,n.fecha_fin_orden,
                   n.tiempo_planeado_min,n.hh_estimada,n.origen,n.detalle_origen,
                   n.programacion_origen_id,n.semana_origen_inicio,n.semana_origen_fin,
                   n.primera_seguimiento_id,n.ultima_seguimiento_id,
                   n.primera_deteccion_en,n.ultima_deteccion_en,n.veces_detectada,
                   CASE
                     WHEN CAST(:latest_id AS bigint) IS NOT NULL
                      AND n.primera_seguimiento_id=CAST(:latest_id AS bigint)
                     THEN true ELSE false
                   END AS nueva_ultima_carga
            FROM programacion.seguimiento_no_programado_v2 n
            WHERE n.programacion_id=:id
            ORDER BY nueva_ultima_carga DESC,n.fecha_fin_orden DESC NULLS LAST,
                     n.primera_deteccion_en DESC,n.numero_ot
        """), {"id": programming_id, "latest_id": latest_id}).mappings()]

        unplanned_new_by_tracking = {
            int(row["tracking_id"]): {
                "cantidad": int(row["cantidad"] or 0),
                "hh": float(row["hh"] or 0),
            }
            for row in conn.execute(text("""
                SELECT primera_seguimiento_id AS tracking_id,
                       count(*)::int AS cantidad,
                       COALESCE(sum(hh_estimada),0) AS hh
                FROM programacion.seguimiento_no_programado_v2
                WHERE programacion_id=:id
                  AND primera_seguimiento_id IS NOT NULL
                GROUP BY primera_seguimiento_id
            """), {"id": programming_id}).mappings()
        }

    for row in history:
        for key in ("hh_programadas", "hh_finalizadas", "hh_pendientes",
                    "avance_ot_pct", "avance_hh_pct"):
            row[key] = float(row[key] or 0)
        daily_unplanned = unplanned_new_by_tracking.get(int(row["id"]), {})
        row["no_programadas_nuevas"] = int(daily_unplanned.get("cantidad") or 0)
        row["hh_no_programadas_nuevas"] = round(float(daily_unplanned.get("hh") or 0), 2)

    by_origin: dict[str, int] = defaultdict(int)
    for row in unplanned_rows:
        row["hh_estimada"] = float(row.get("hh_estimada") or 0)
        row["tiempo_planeado_min"] = (
            float(row["tiempo_planeado_min"]) if row.get("tiempo_planeado_min") is not None else None
        )
        by_origin[str(row.get("origen") or "SIN_CLASIFICAR")] += 1

    unplanned_total = len(unplanned_rows)
    unplanned_new = sum(1 for row in unplanned_rows if row.get("nueva_ultima_carga"))
    unplanned_hh = round(sum(float(row.get("hh_estimada") or 0) for row in unplanned_rows), 2)
    programmed_finalized = int(history[0]["finalizados"] or 0) if history else 0
    execution_total = programmed_finalized + unplanned_total

    return {
        "programming": dict(programming),
        "latest": history[0] if history else None,
        "history": history,
        "by_area": _area_summary(latest_rows),
        "rows": latest_rows,
        "unplanned": {
            "summary": {
                "total": unplanned_total,
                "nuevas_ultima_carga": unplanned_new,
                "hh_estimada": unplanned_hh,
                "ejecucion_total": execution_total,
                "participacion_pct": _pct(unplanned_total, execution_total),
            },
            "by_origin": [
                {"origen": origin, "cantidad": count}
                for origin, count in sorted(by_origin.items())
            ],
            "rows": unplanned_rows,
        },
    }

def record_week_tracking(
    *,
    programming_id: int,
    content: bytes,
    filename: str,
    recorded_by: str | None = None,
) -> dict[str, Any]:
    try:
        preview = preview_week_closure(programming_id=programming_id, content=content)
        calendar_rows = _parse_calendar(content)
    except V2ClosureError as exc:
        raise V2WeeklyTrackingError(str(exc)) from exc

    if preview["programming"]["estado"] == "CERRADA":
        raise V2WeeklyTrackingError(
            "La semana ya está cerrada. El seguimiento histórico se conserva en modo consulta."
        )

    rows = [dict(row) for row in preview["rows"]]

    # Seguimiento acumulativo: una OT ya detectada como finalizada no vuelve a
    # aparecer como pendiente/sin coincidencia en una carga posterior.
    with get_engine().connect() as conn:
        previous_finalized = {
            int(row[0])
            for row in conn.execute(text("""
                SELECT DISTINCT si.orden_mantenimiento_id
                FROM programacion.seguimiento_semanal_item_v2 si
                JOIN programacion.seguimiento_semanal_v2 s ON s.id=si.seguimiento_id
                WHERE s.programacion_id=:id
                  AND si.finalizado=true
            """), {"id": programming_id}).all()
        }

    for row in rows:
        if int(row["orden_mantenimiento_id"]) in previous_finalized and row.get("finalizado") is not True:
            row["finalizado"] = True
            row["estado_excel"] = row.get("estado_excel") or "FINALIZADA PREVIAMENTE"
            row["coincidencia"] = "SEGUIMIENTO PREVIO"

    programmed = len(rows)
    finalized = sum(1 for row in rows if row.get("finalizado") is True)
    pending = sum(1 for row in rows if row.get("finalizado") is False)
    unmatched = sum(1 for row in rows if row.get("finalizado") is None)
    hh_programmed = round(sum(float(row.get("hh") or 0) for row in rows), 2)
    hh_finalized = round(sum(float(row.get("hh") or 0) for row in rows if row.get("finalizado") is True), 2)
    hh_pending = round(sum(float(row.get("hh") or 0) for row in rows if row.get("finalizado") is False), 2)
    ot_pct = _pct(finalized, programmed)
    hh_pct = _pct(hh_finalized, hh_programmed)
    unplanned_detected = 0
    unplanned_new = 0
    backlog_reconciliation = {
        "revisadas": 0,
        "encontradas": 0,
        "finalizadas": 0,
        "abiertas": 0,
        "ambiguas": 0,
        "finalizadas_ots": [],
    }

    with get_engine().begin() as conn:
        tracking_id = conn.execute(text("""
            INSERT INTO programacion.seguimiento_semanal_v2(
              programacion_id,archivo_nombre,registrado_por,
              total_items,finalizados,pendientes,sin_coincidencia,
              hh_programadas,hh_finalizadas,hh_pendientes,
              avance_ot_pct,avance_hh_pct
            ) VALUES(
              :programming_id,:filename,:recorded_by,
              :programmed,:finalized,:pending,:unmatched,
              :hh_programmed,:hh_finalized,:hh_pending,
              :ot_pct,:hh_pct
            )
            RETURNING id
        """), {
            "programming_id": programming_id,
            "filename": filename,
            "recorded_by": recorded_by,
            "programmed": programmed,
            "finalized": finalized,
            "pending": pending,
            "unmatched": unmatched,
            "hh_programmed": hh_programmed,
            "hh_finalized": hh_finalized,
            "hh_pending": hh_pending,
            "ot_pct": ot_pct,
            "hh_pct": hh_pct,
        }).scalar_one()

        if rows:
            payload = [{
                "tracking_id": tracking_id,
                "program_item_id": int(row["programacion_item_id"]),
                "order_id": int(row["orden_mantenimiento_id"]),
                "numero_ot": row.get("numero_ot"),
                "area_codigo": row.get("area_codigo"),
                "area_nombre": row.get("area_nombre"),
                "hh": float(row.get("hh") or 0),
                "state": row.get("estado_excel"),
                "finalized": row.get("finalizado"),
                "match": row.get("coincidencia"),
            } for row in rows]
            conn.execute(text("""
                INSERT INTO programacion.seguimiento_semanal_item_v2(
                  seguimiento_id,programacion_item_id,orden_mantenimiento_id,
                  numero_ot,area_codigo,area_nombre,hh_programadas,
                  estado_calendario,finalizado,coincidencia
                ) VALUES(
                  :tracking_id,:program_item_id,:order_id,
                  :numero_ot,:area_codigo,:area_nombre,:hh,
                  :state,:finalized,:match
                )
            """), payload)

        backlog_reconciliation = _reconcile_backlog_from_calendar(
            conn,
            specialty=str(preview["programming"].get("especialidad") or ""),
            calendar_rows=calendar_rows,
            recorded_by=recorded_by,
        )

        unplanned_rows = _detect_unplanned_rows(
            conn,
            programming=dict(preview["programming"]),
            programmed_rows=rows,
            calendar_rows=calendar_rows,
        )
        unplanned_detected = len(unplanned_rows)

        if unplanned_rows:
            existing_ots = {
                _ot_key(row[0])
                for row in conn.execute(text("""
                    SELECT numero_ot
                    FROM programacion.seguimiento_no_programado_v2
                    WHERE programacion_id=:id
                """), {"id": programming_id}).all()
            }
            unplanned_new = sum(
                1 for row in unplanned_rows if _ot_key(row["numero_ot"]) not in existing_ots
            )
            unplanned_payload = [{
                **row,
                "numero_ot": _ot_key(row["numero_ot"]),
                "tracking_id": int(tracking_id),
            } for row in unplanned_rows]
            conn.execute(text("""
                INSERT INTO programacion.seguimiento_no_programado_v2(
                  programacion_id,orden_mantenimiento_id,numero_ot,
                  activo_codigo,activo_descripcion,area_codigo,area_nombre,
                  plan_trabajo,especialidad,estado_calendario,fecha_fin_orden,
                  tiempo_planeado_min,hh_estimada,origen,detalle_origen,
                  programacion_origen_id,semana_origen_inicio,semana_origen_fin,
                  primera_seguimiento_id,ultima_seguimiento_id
                ) VALUES(
                  :programacion_id,:orden_mantenimiento_id,:numero_ot,
                  :activo_codigo,:activo_descripcion,:area_codigo,:area_nombre,
                  :plan_trabajo,:especialidad,:estado_calendario,:fecha_fin_orden,
                  :tiempo_planeado_min,:hh_estimada,:origen,:detalle_origen,
                  :programacion_origen_id,:semana_origen_inicio,:semana_origen_fin,
                  :tracking_id,:tracking_id
                )
                ON CONFLICT(programacion_id,numero_ot) DO UPDATE SET
                  orden_mantenimiento_id=COALESCE(EXCLUDED.orden_mantenimiento_id,
                    programacion.seguimiento_no_programado_v2.orden_mantenimiento_id),
                  activo_codigo=COALESCE(EXCLUDED.activo_codigo,
                    programacion.seguimiento_no_programado_v2.activo_codigo),
                  activo_descripcion=COALESCE(EXCLUDED.activo_descripcion,
                    programacion.seguimiento_no_programado_v2.activo_descripcion),
                  area_codigo=COALESCE(EXCLUDED.area_codigo,
                    programacion.seguimiento_no_programado_v2.area_codigo),
                  area_nombre=COALESCE(EXCLUDED.area_nombre,
                    programacion.seguimiento_no_programado_v2.area_nombre),
                  plan_trabajo=COALESCE(EXCLUDED.plan_trabajo,
                    programacion.seguimiento_no_programado_v2.plan_trabajo),
                  especialidad=COALESCE(EXCLUDED.especialidad,
                    programacion.seguimiento_no_programado_v2.especialidad),
                  estado_calendario=EXCLUDED.estado_calendario,
                  fecha_fin_orden=EXCLUDED.fecha_fin_orden,
                  tiempo_planeado_min=COALESCE(EXCLUDED.tiempo_planeado_min,
                    programacion.seguimiento_no_programado_v2.tiempo_planeado_min),
                  hh_estimada=EXCLUDED.hh_estimada,
                  origen=EXCLUDED.origen,
                  detalle_origen=EXCLUDED.detalle_origen,
                  programacion_origen_id=EXCLUDED.programacion_origen_id,
                  semana_origen_inicio=EXCLUDED.semana_origen_inicio,
                  semana_origen_fin=EXCLUDED.semana_origen_fin,
                  ultima_seguimiento_id=EXCLUDED.ultima_seguimiento_id,
                  ultima_deteccion_en=now(),
                  veces_detectada=programacion.seguimiento_no_programado_v2.veces_detectada+1
            """), unplanned_payload)

    result = get_week_tracking(programming_id)
    result["saved"] = {
        "tracking_id": int(tracking_id),
        "archivo": filename,
        "finalizadas": finalized,
        "pendientes": pending,
        "sin_coincidencia": unmatched,
        "avance_ot_pct": ot_pct,
        "avance_hh_pct": hh_pct,
        "no_programadas_detectadas": unplanned_detected,
        "no_programadas_nuevas": unplanned_new,
        "backlog_revisadas": int(backlog_reconciliation["revisadas"]),
        "backlog_encontradas": int(backlog_reconciliation["encontradas"]),
        "backlog_finalizadas": int(backlog_reconciliation["finalizadas"]),
        "backlog_abiertas": int(backlog_reconciliation["abiertas"]),
        "backlog_ambiguas": int(backlog_reconciliation["ambiguas"]),
        "backlog_finalizadas_ots": list(backlog_reconciliation["finalizadas_ots"]),
    }
    return result

def export_week_tracking_excel(
    programming_id: int,
    area: str | None = None,
) -> tuple[bytes, str]:
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    data = get_week_tracking(programming_id)
    programming = data["programming"]
    rows = list(data.get("rows") or [])
    area = (area or "").strip().upper()
    if area:
        rows = [row for row in rows if str(row.get("area_codigo") or "").upper() == area]

    week_start = programming["semana_inicio"]
    week_end = programming["semana_fin"]
    week_number = _week_number(week_start)
    week_label = f"SEMANA {week_number}" if week_number is not None else "TRANSICIÓN"
    specialty = str(programming["especialidad"] or "").upper()
    specialty_name = {
        "MEC": "MECÁNICA",
        "ELE": "ELÉCTRICA",
        "MET": "METROLOGÍA",
        "SER": "SERVICIOS",
    }.get(specialty, specialty)

    wb = Workbook()
    ws = wb.active
    ws.title = (f"{specialty} - {week_label.title()}")[:31]
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A8"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    green = "009B5A"
    dark_green = "16603D"
    light_green = "EAF7EF"
    yellow = "F4C20D"
    pale_yellow = "FFF8D9"
    navy = "17365D"
    gray = "67776F"
    light_gray = "F2F4F3"
    line = Side(style="thin", color="D6E4DC")
    border = Border(left=line, right=line, top=line, bottom=line)

    try:
        logo_path = Path(__file__).resolve().parents[1] / "assets" / "team_foods_logo.b64"
        logo_bytes = base64.b64decode(logo_path.read_text(encoding="utf-8").strip())
        image = XLImage(BytesIO(logo_bytes))
        image.width = 145
        image.height = 56
        ws.add_image(image, "A1")
    except Exception:
        ws["A1"] = "TEAM FOODS"
        ws["A1"].font = Font(size=16, bold=True, color=green)

    ws.merge_cells("C1:J2")
    ws["C1"] = f"SEGUIMIENTO DE PROGRAMACIÓN · {week_label}"
    ws["C1"].font = Font(size=18, bold=True, color=dark_green)
    ws["C1"].alignment = Alignment(vertical="center")

    ws.merge_cells("C3:J3")
    area_text = f" · ÁREA {area}" if area else " · TODAS LAS ÁREAS"
    ws["C3"] = (
        f"Planta Barranquilla · {specialty_name}{area_text} · "
        f"{week_start:%d/%m/%Y} al {week_end:%d/%m/%Y}"
    )
    ws["C3"].font = Font(size=9, bold=True, color=gray)

    ws.merge_cells("K1:L1")
    ws["K1"] = "CEK GLOBAL"
    ws["K1"].font = Font(size=12, bold=True, color=navy)
    ws["K1"].alignment = Alignment(horizontal="right")
    ws.merge_cells("K2:L2")
    ws["K2"] = "Inspection Services"
    ws["K2"].font = Font(size=8, bold=True, color=gray)
    ws["K2"].alignment = Alignment(horizontal="right")

    total = len(rows)
    finalized = sum(1 for row in rows if row.get("finalizado") is True)
    pending = sum(1 for row in rows if row.get("finalizado") is False)
    unmatched = sum(1 for row in rows if row.get("finalizado") is None)
    progress = _pct(finalized, total)

    cards = [
        ("A5:C5", "PROGRAMADAS", total, green, light_green),
        ("D5:F5", "FINALIZADAS", finalized, green, light_green),
        ("G5:I5", "PENDIENTES", pending, yellow, pale_yellow),
        ("J5:L5", "AVANCE OT", progress, green, light_green),
    ]
    for cell_range, label, value, accent, fill in cards:
        ws.merge_cells(cell_range)
        start = cell_range.split(":")[0]
        cell = ws[start]
        cell.value = f"{label}  ·  {value:.1f}%" if label == "AVANCE OT" else f"{label}  ·  {value}"
        cell.font = Font(size=10, bold=True, color=dark_green if accent != yellow else navy)
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        for row_cells in ws[cell_range]:
            for item in row_cells:
                item.border = border

    ws.merge_cells("A6:L6")
    ws["A6"] = (
        "Las filas tachadas corresponden a OT detectadas como FINALIZADAS en el seguimiento. "
        "El seguimiento no sustituye el Cierre semanal oficial."
    )
    ws["A6"].font = Font(size=8, italic=True, color=gray)
    ws["A6"].alignment = Alignment(vertical="center")

    headers = [
        "Estado", "OT", "Área", "Equipo", "Descripción equipo", "Plan de trabajo",
        "Condición", "Origen", "Personas", "Tiempo min", "H-H", "Estado calendario",
    ]
    for col, label in enumerate(headers, 1):
        cell = ws.cell(7, col, label)
        cell.font = Font(size=8, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=dark_green)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    for row_number, row in enumerate(rows, 8):
        finalized_row = row.get("finalizado") is True
        state = (
            "FINALIZADA" if finalized_row
            else "PENDIENTE" if row.get("finalizado") is False
            else "SIN VALIDAR"
        )
        origin = "BACKLOG" if row.get("origen") == "BACKLOG" or row.get("origen_backlog") else "PMP DEL MES"
        values = [
            state,
            row.get("numero_ot") or "SIN ASIGNAR",
            row.get("area_codigo") or "",
            row.get("activo_codigo") or "",
            row.get("activo_descripcion") or "",
            row.get("plan_trabajo") or "",
            "EQUIPO DETENIDO" if row.get("requiere_parada") else "EQUIPO OPERANDO",
            origin,
            row.get("personas") if row.get("personas") is not None else "",
            row.get("tiempo_min") if row.get("tiempo_min") is not None else "",
            float(row.get("hh_programadas") or 0),
            row.get("estado_calendario") or "",
        ]
        for col, value in enumerate(values, 1):
            cell = ws.cell(row_number, col, value)
            cell.border = border
            cell.alignment = Alignment(
                vertical="top",
                horizontal="center" if col in (1, 3, 7, 8, 9, 10, 11) else "left",
                wrap_text=True,
            )
            if finalized_row:
                cell.font = Font(size=8, strike=True, color="74827A")
                cell.fill = PatternFill("solid", fgColor=light_gray)
            else:
                cell.font = Font(size=8, color="20382E")
                if state == "PENDIENTE":
                    cell.fill = PatternFill("solid", fgColor="FFFCF0")
        if finalized_row:
            ws.cell(row_number, 1).font = Font(size=8, bold=True, strike=True, color=dark_green)
        elif state == "PENDIENTE":
            ws.cell(row_number, 1).font = Font(size=8, bold=True, color="8A6800")

    widths = [14, 17, 12, 21, 31, 39, 20, 15, 10, 12, 10, 22]
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(idx)].width = width

    last_row = max(7, 7 + len(rows))
    ws.auto_filter.ref = f"A7:L{last_row}"
    ws.print_title_rows = "1:7"
    ws.print_area = f"A1:L{last_row}"
    ws.oddFooter.left.text = "Team Foods · CEK Global Inspection Services"
    ws.oddFooter.right.text = "Página &P de &N"

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    week_file = f"semana_{week_number}" if week_number is not None else "transicion"
    area_file = f"_{area}" if area else ""
    return output.getvalue(), f"seguimiento_{specialty}_{week_file}{area_file}.xlsx"
