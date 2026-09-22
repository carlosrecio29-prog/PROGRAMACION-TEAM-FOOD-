"""Contract for temporary calendar preview before the next monthly programming."""
from datetime import date
from io import BytesIO

from openpyxl import Workbook, load_workbook

from backend.services.v2_advance_stops import (
    _read_calendar, classify_advance, export_advance_excel,
)


def sample_rows():
    return [
        {"fila_origen":2,"activo_codigo":"BA-1","plan_clave_software":"44-RUTINA MEC",
         "numero_ot_raw":"","especialidad":"MEC","tiempo_planeado_min":60,
         "cronograma_planeacion":"SEPTIEMBRE","estado":None},
        {"fila_origen":3,"activo_codigo":"BA-1","plan_clave_software":"45-INSPECCION",
         "numero_ot_raw":"OT-10","especialidad":"MEC","tiempo_planeado_min":30,
         "cronograma_planeacion":"","estado":"ABIERTO"},
        {"fila_origen":4,"activo_codigo":"BA-2","plan_clave_software":"46-SIN TIEMPO",
         "numero_ot_raw":"","especialidad":"ELE","tiempo_planeado_min":120,
         "cronograma_planeacion":"","estado":None},
        {"fila_origen":5,"activo_codigo":"BA-1","plan_clave_software":"47-OPERACION",
         "numero_ot_raw":"","especialidad":"MEC","tiempo_planeado_min":10,
         "cronograma_planeacion":"","estado":None},
        {"fila_origen":6,"activo_codigo":"BA-99","plan_clave_software":"48-NO EXISTE",
         "numero_ot_raw":"","especialidad":"ELE","tiempo_planeado_min":10,
         "cronograma_planeacion":"","estado":None},
    ]


def sample_plans():
    def plan(group, name, stop, *, operacion=False):
        return {"grupo":group, "plan_trabajo":name,
                "tiempo_parada_efectivo_min":stop,"es_operacion":operacion,
                "descripcion_plan_trabajo":name, "especialidad":"MEC",
                "numero_personas_efectivo":2,"tiempo_ejecucion_min":60}
    return {
        "44-RUTINA MEC": plan("44","RUTINA MEC",60),
        "45-INSPECCION": plan("45","INSPECCION",0),
        "46-SIN TIEMPO": plan("46","SIN TIEMPO",None),
        "47-OPERACION": plan("47","OPERACION",60,operacion=True),
    }


def sample_assets():
    return {
        "BA-1":{"codigo":"BA-1","descripcion":"Equipo uno",
                "area_codigo":"JB","criticidad":"A"},
        "BA-2":{"codigo":"BA-2","descripcion":"Equipo dos",
                "area_codigo":"FR","criticidad":"B"},
    }


def test_no_ot_does_not_hide_downtime_and_zero_means_operating():
    result=classify_advance(sample_rows(),sample_plans(),
                            sample_assets(),date(2026,10,1))
    assert result["periodo"]=="2026-10-01"
    assert result["total_archivo"]==5
    assert result["excluidos_operacion"]==1
    assert result["equipo_detenido"]==1
    assert result["operando"]==1
    assert result["sin_definir"]==2
    assert result["sin_ot"]==3
    stopped=[r for r in result["filas"] if r["condicion"]=="EQUIPO DETENIDO"]
    assert stopped[0]["numero_ot"]==""
    assert stopped[0]["tiempo_parada_min"]==60
    assert stopped[0]["hh_estimadas"]==2
    operating=[r for r in result["filas"] if r["condicion"]=="OPERANDO"]
    assert operating[0]["tiempo_parada_min"]==0


def test_export_separates_conditions_and_includes_unassigned_ot():
    result=classify_advance(sample_rows(),sample_plans(),
                            sample_assets(),date(2026,10,1))
    payload,name=export_advance_excel(result)
    assert payload.startswith(b"PK")
    assert name=="anticipacion_paradas_202610_team_foods_cek.xlsx"
    book=load_workbook(BytesIO(payload),read_only=True)
    assert book.sheetnames==["RESUMEN","EQUIPO DETENIDO","OPERANDO","SIN DEFINIR"]
    assert book["EQUIPO DETENIDO"]["B5"].value is None
    assert book["OPERANDO"]["B5"].value=="OT-10"
    assert book["SIN DEFINIR"].max_row==6
    assert book["RESUMEN"]["B8"].value==1
    book.close()


def test_accepts_excel_without_orden_header():
    book=Workbook()
    sheet=book.active
    sheet.append(["Activo","PlanTrabajo","Especialidad","TiempoPlaneado"])
    sheet.append(["BA-1","44-RUTINA MEC","MEC",60])
    stream=BytesIO()
    book.save(stream)
    rows=_read_calendar(stream.getvalue())
    assert len(rows)==1
    assert rows[0]["numero_ot_raw"]==""


def test_rejects_irrelevant_excel_not_empty_preview():
    import pytest
    book=Workbook()
    sheet=book.active
    sheet.append(["Otra columna"])
    sheet.append(["OT-100"])
    stream=BytesIO()
    book.save(stream)
    with pytest.raises(ValueError,match="Activo y PlanTrabajo"):
        _read_calendar(stream.getvalue())


def test_no_write_sql_in_preview_service():
    from pathlib import Path
    source=(Path(__file__).resolve().parents[1]
            /"backend/services/v2_advance_stops.py").read_text()
    assert "get_engine().connect()" in source
    for word in ("INSERT INTO programacion.", "UPDATE programacion.",
                 "DELETE FROM programacion."):
        assert word not in source


def test_plan_spacing_variation_matches_and_operation_is_excluded():
    plans = sample_plans()
    plans["27-OPERACION - RUTINA ELECTRICA"] = {
        "grupo":"27", "plan_trabajo":"OPERACION - RUTINA ELECTRICA",
        "es_operacion":True, "tiempo_parada_efectivo_min":None,
    }
    entries = [
        {**sample_rows()[0], "plan_clave_software": "44 - RUTINA MEC"},
        {**sample_rows()[0], "fila_origen":8,
         "plan_clave_software": "27-OPERACION -RUTINA ELECTRICA"},
        {**sample_rows()[0], "fila_origen":9,
         "plan_clave_software": "28-OPERACION - RUTINA NO ESTA EN MAESTRO"},
    ]
    r=classify_advance(entries,plans,sample_assets(),date(2026,10,1))
    assert r["equipo_detenido"] == 1
    assert r["excluidos_operacion"] == 2
    assert r["sin_definir"] == 0


def test_distinguishes_missing_plan_from_missing_downtime():
    r=classify_advance(sample_rows(),sample_plans(),sample_assets(),date(2026,10,1))
    assert r["sin_definir_por_tiempo"] == 1
    assert r["sin_definir_por_plan"] == 1
    reasons={row["motivo_sin_definir"] for row in r["filas"]}
    assert "PLAN NO ENCONTRADO" in reasons
    assert "TIEMPO PARADA VACÍO" in reasons


def test_calendar_uses_description_alias_from_uploaded_master():
    plans = {
        "200-MANTENIMIENTO": {
            "grupo": "200", "plan_trabajo": "MANTENIMIENTO",
            "descripcion_plan_trabajo": "Mantenimiento y Limpieza de Equipos Retráctiles",
            "es_operacion": False, "tiempo_parada_efectivo_min": 20,
            "numero_personas_efectivo": 1, "tiempo_ejecucion_min": 20,
            "especialidad": "MEC",
        },
        "109-COMPARACION MEDIDA CON PATRON PARA SENSOR DE TEMPERATURA SEMESTRAL": {
            "grupo": "109",
            "plan_trabajo": "COMPARACION MEDIDA CON PATRON PARA SENSOR DE TEMPERATURA SEMESTRAL",
            "descripcion_plan_trabajo": "COMPARACION MEDIDA CON PATRON PARA RTD SEMESTRAL",
            "es_operacion": False, "tiempo_parada_efectivo_min": 60,
            "numero_personas_efectivo": 1, "tiempo_ejecucion_min": 60,
            "especialidad": "MET",
        }
    }
    entries=[
        {**sample_rows()[0], "plan_clave_software":
            "200-Mantenimiento y Limpieza de Equipos Retráctiles"},
        {**sample_rows()[0], "fila_origen": 3, "plan_clave_software":
            "109-COMPARACION MEDIDA CON PATRON PARA RTD SEMESTRAL"},
    ]
    result=classify_advance(entries,plans,sample_assets(),date(2026,10,1))
    assert result["equipo_detenido"] == 2
    assert result["sin_definir_por_plan"] == 0
    assert {r["tiempo_parada_min"] for r in result["filas"]} == {20,60}


def test_description_alias_never_guesses_ambiguous_plan():
    plans={
        "90-PLAN A": {"grupo": "90", "plan_trabajo": "PLAN A",
                      "descripcion_plan_trabajo": "RUTINA GENERICA",
                      "es_operacion": False,"tiempo_parada_efectivo_min": 0,
                      "numero_personas_efectivo":1,"tiempo_ejecucion_min":15,
                      "especialidad":"ELE"},
        "90-PLAN B": {"grupo": "90", "plan_trabajo": "PLAN B",
                      "descripcion_plan_trabajo": "RUTINA GENERICA",
                      "es_operacion": False,"tiempo_parada_efectivo_min": 90,
                      "numero_personas_efectivo":1,"tiempo_ejecucion_min":90,
                      "especialidad":"ELE"},
    }
    entries=[{**sample_rows()[0], "plan_clave_software":"90-RUTINA GENERICA"}]
    result=classify_advance(entries,plans,sample_assets(),date(2026,10,1))
    assert result["sin_definir_por_ambiguo"]==1
    assert result["equipo_detenido"]==0
    assert result["operando"]==0
